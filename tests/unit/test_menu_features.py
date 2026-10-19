import base64
from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from homeplan.checks import evaluate
from homeplan.importer import LABELS, LISTING_COLS
from homeplan.models import LoanPrepayment
from homeplan.server import changes_between, create_app


def make_excel(*, income: int = 4500000, cash: int = 120000000) -> bytes:
    """개인 파일을 사용하지 않는 작은 v18 형식 가상 입력."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ["기본설정", "조건별비교", "월별가계흐름", "우리자금", "아파트매물", "결과요약"]:
        workbook.create_sheet(name)
    settings = {
        "asset_base_date": "2026-10-19",
        "contract_date": "2026-11-20",
        "closing_date": "2027-02-26",
        "move_in_date": "2027-03-01",
        "husband_net_pay": income,
    }
    for row, (key, value) in enumerate(settings.items(), 1):
        workbook["기본설정"].cell(row, 2, LABELS[key])
        workbook["기본설정"].cell(row, 5, value)
    workbook["조건별비교"]["B5"] = 84.9
    for col, (key, label) in enumerate(LISTING_COLS.items(), 1):
        workbook["아파트매물"].cell(5, col, label)
        value = {
            "id": "EXCEL-HOME",
            "data_kind": "실제",
            "name": "가상 매물",
            "area_m2": 84.9,
            "price": 500000000,
            "ask_price": 500000000,
            "tenant": "집주인 거주",
        }.get(key)
        if value is not None:
            workbook["아파트매물"].cell(6, col, value)
    workbook["결과요약"]["C6"] = "EXCEL-HOME"
    for col, value in {
        1: "EXCEL-CASH",
        2: "실제",
        3: "공동",
        4: "현금·예금",
        5: "가상 예금",
        6: cash,
        10: "2026-10-19",
        11: True,
        12: "확정",
    }.items():
        workbook["우리자금"].cell(6, col, value)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.fixture
def menu_client(tmp_path, household):
    (tmp_path / "data").mkdir()
    (tmp_path / "data/demo.local.json").write_text(household.model_dump_json(), "utf-8")
    with TestClient(create_app(tmp_path / "qa.sqlite", seed_root=tmp_path)) as client:
        yield client


def test_excel_preview_does_not_save_and_merge_preserves_app_edits(menu_client, household):
    client = menu_client
    before = client.get("/api/workspaces").json()
    payload = {"filename": "synthetic.xlsx", "content": base64.b64encode(make_excel()).decode()}
    response = client.post("/api/import-excel/preview", json=payload)
    assert response.status_code == 200, response.text
    imported = response.json()["household"]
    assert imported["settings"]["husband_net_pay"] == 4500000
    assert imported["selected_listing"] == "EXCEL-HOME"
    assert imported["_source"] == "synthetic.xlsx"
    assert client.get("/api/workspaces").json() == before
    existing = household.model_dump(mode="json")
    existing["app_setting_keys"] = ["husband_net_pay"]
    existing["settings"]["husband_net_pay"] = 6100000
    existing["fund_events"].append(
        {**imported["fund_events"][0], "value": 9000000, "origin": "app"}
    )
    result = client.post(
        "/api/import-excel/preview",
        json={**payload, "household": existing, "refresh_imported": True},
    )
    assert result.status_code == 200, result.text
    merged = result.json()["household"]
    assert merged["settings"]["husband_net_pay"] == 6100000
    assert next(e for e in merged["fund_events"] if e["id"] == "EXCEL-CASH")["value"] == 9000000
    assert len(merged["listings"]) == len(existing["listings"]) + 1
    assert client.get("/api/workspaces").json() == before


@pytest.mark.parametrize("filename,content", [("bad.txt", b"bad"), ("bad.xlsx", b"bad")])
def test_excel_preview_reports_invalid_files(menu_client, filename, content):
    client = menu_client
    response = client.post(
        "/api/import-excel/preview",
        json={
            "filename": filename,
            "content": base64.b64encode(content).decode(),
        },
    )
    assert response.status_code == 422
    assert isinstance(response.json()["detail"], str)


def test_repayment_details_conserve_principal_with_same_day_extra_payment(household, rules):
    initial = evaluate(household, None, rules, analyze_assumptions=False)
    day = initial.loans["mortgage_schedule"][0]["date"]
    household.loan_prepayments = [
        LoanPrepayment(
            id="EXTRA",
            loan="mortgage",
            date=day,
            amount=10000000,
            fee=1234,
        )
    ]
    result = evaluate(household, None, rules, analyze_assumptions=False)
    rows = result.loans["repayment_details"]["mortgage"]
    assert [r["kind"] for r in rows[:2]] == ["정기 납부", "추가 상환"]
    assert rows[1]["total"] == 10001234
    assert rows[1]["balance"] == rows[0]["balance"] - 10000000
    assert sum(r["principal"] for r in rows) == result.loans["mortgage"]
    assert rows[-1]["balance"] == 0
    assert sum(r["interest"] for r in rows) < sum(
        r["interest"] for r in initial.loans["mortgage_schedule"]
    )


def test_bank_quote_changes_are_in_save_preview(household):
    before = household.model_dump(mode="json")
    after = household.model_dump(mode="json")
    after["bank_quotes"] = [
        {"id": "BANK1", "bank": "가상은행", "date": date(2026, 10, 19).isoformat(), "options": []}
    ]
    changes = changes_between(before, after)
    assert any(c["kind"] == "bank_quotes" and c["key"] == "BANK1" for c in changes)


def test_excel_merge_remains_refreshable_after_save_but_later_app_edit_is_preserved(menu_client):
    client = menu_client
    stored = client.get("/api/workspaces/demo").json()

    def preview(existing, income, cash):
        response = client.post(
            "/api/import-excel/preview",
            json={
                "filename": "synthetic.xlsx",
                "content": base64.b64encode(make_excel(income=income, cash=cash)).decode(),
                "household": existing,
                "refresh_imported": True,
            },
        )
        assert response.status_code == 200, response.text
        return response.json()["household"]

    candidate = preview(stored["household"], 4500000, 120000000)
    saved = client.put(
        "/api/workspaces/demo", json={"household": candidate, "revision": stored["revision"]}
    ).json()
    assert "_excel_import_baseline" not in saved["household"]
    assert "husband_net_pay" not in saved["household"]["app_setting_keys"]
    assert (
        next(e for e in saved["household"]["fund_events"] if e["id"] == "EXCEL-CASH")["origin"]
        == "excel"
    )
    candidate = preview(saved["household"], 4800000, 130000000)
    assert candidate["settings"]["husband_net_pay"] == 4800000
    # User changes an imported value after the preview, before saving.
    candidate["settings"]["husband_net_pay"] = 6200000
    saved = client.put(
        "/api/workspaces/demo", json={"household": candidate, "revision": saved["revision"]}
    ).json()
    assert "husband_net_pay" in saved["household"]["app_setting_keys"]
    candidate = preview(saved["household"], 4900000, 140000000)
    assert candidate["settings"]["husband_net_pay"] == 6200000
    assert (
        next(e for e in candidate["fund_events"] if e["id"] == "EXCEL-CASH")["value"] == 140000000
    )
