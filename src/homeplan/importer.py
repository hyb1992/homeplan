"""엑셀 I/O: 알려진 날짜 수식 보존, 수동 상담·앱 입력 병합."""

import re
import warnings
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .dates import ALIASES, DATE_KEYS
from .storage import read_json, save_json

LABELS = {
    # 일정
    "asset_base_date": "자산 기준일",
    "contract_date": "계약일",
    "deposit_ratio": "계약금 비율",
    "interim_date": "중도금 지급일(있는 경우)",
    "interim_amount": "중도금 금액",
    "closing_date": "잔금일(대출 실행일)",
    "move_in_date": "입주일",
    "lease_end_date": "기존 전세계약 만기(참고)",
    "lease_refund_date": "기존 전세보증금 반환일",
    "lease_loan_repay_date": "기존 전세대출 상환일",
    "same_day_inflow_usable": "잔금·계약 당일 들어오는 돈을 그날 쓸 수 있는가",
    # 출산·휴직
    "birth_date": "출산 예정일(가정일)",
    "maternity_start": "아내 출산휴가 시작일",
    "maternity_months": "출산휴가 개월 수",
    "maternity_pay_m12": "출산휴가 1~2개월차 월 수령액(실수령)",
    "maternity_pay_m3": "출산휴가 3개월차 월 수령액",
    "wife_leave_start": "아내 육아휴직 시작일",
    "wife_leave_months": "아내 육아휴직 개월 수(쉬는 기간)",
    "return_delay_months": "복직 지연(개월)",
    "wife_gov_benefit_months": "정부 육아휴직급여 지급 개월 수(아내)",
    "wife_ordinary_wage": "아내 통상임금(월, 세전)",
    "wife_company_topup": "아내 회사 보전금(월)",
    "wife_company_topup_months": "아내 회사 보전 개월 수",
    "benefit_pay_delay_months": "정부급여 지급 지연(개월)",
    "husband_leave_use": "남편 육아휴직 사용",
    "husband_leave_start": "남편 육아휴직 시작일",
    "husband_leave_months": "남편 육아휴직 개월 수",
    "husband_ordinary_wage": "남편 통상임금(월, 세전)",
    "husband_gov_benefit_months": "정부 육아휴직급여 지급 개월 수(남편)",
    "parents_6p6": "부모 함께 육아휴직 특례(6+6) 상한 적용",
    # 소득·지출
    "husband_net_pay": "남편 월 실수령액",
    "wife_net_pay": "아내 월 실수령액",
    "payday": "급여일(매월 며칠)",
    "living_cost": "월 생활비",
    "mgmt_fee_default": "입주 후 월 관리비·공과금(기본값)",
    "mgmt_in_living": "  └ 관리비가 월 생활비에 이미 포함?",
    "fixed_cost": "보험·차량·통신 등 월 고정비",
    "annual_cost": "연 1회 비용(자동차보험 등)",
    "annual_cost_month": "  └ 연 1회 비용 납부월",
    "fixed_in_living": "  └ 위 고정비·연 비용이 월 생활비에 이미 포함?",
    "childcare_cost": "월 양육비(출산월부터)",
    "birth_one_off": "출산 관련 일회성 비용",
    "property_tax_annual": "연간 보유세(재산세 등) 예상액",
    "property_tax_month1": "  └ 보유세 납부월 ①",
    "property_tax_month2": "  └ 보유세 납부월 ②",
    "lease_loan_rate": "기존 전세대출 금리(연)",
    "lease_interest_in_living": "  └ 전세대출 이자가 월 생활비에 이미 포함?",
    "monthly_saving": "월 적금·투자 납입액",
    "saving_end_date": "  └ 납입 종료일",
    # 한도
    "monthly_repay_limit": "월 대출 상환 한도",
    "monthly_spend_limit": "월 가계 전체 지출 한도",
    "min_emergency_cash": "최소 비상금(전 기간 유지할 현금)",
    "target_cash_after_purchase": "매수 직후 남길 목표 현금",
    # 주담대
    "mortgage_mode": "대출 모드",
    "mortgage_desired": "기본 대출금",
    "mortgage_cap": "대출 상한(사용자 설정)",
    "mortgage_rate": "대출 금리(연, 실제 납부 기준)",
    "mortgage_rate_type": "금리 유형",
    "mortgage_reset_after_months": "금리 재산정 시점(대출 실행 후 개월)",
    "mortgage_reset_rate": "재산정 후 금리(연)",
    "mortgage_years": "상환기간(년)",
    "mortgage_method": "상환 방식",
    "equal_principal_dsr_basis": "원금균등 DSR 산정 방식",
    "ltv": "LTV 적용률",
    "dsr_limit": "DSR 한도",
    "stress_rate": "스트레스 가산금리(심사용)",
    "bank_income": "은행 인정 연소득(세전)",
    "existing_debt_annual": "기존 부채 연 원리금(대출 실행 시점)",
    "bank_confirmed_limit": "은행 확인 대출 한도(있으면 추정값 대신 적용)",
    "loan_status": "대출 확인 상태",
    # 신용대출·청담대·가족차입
    "husband_gross": "남편 총급여(전년도 원천징수영수증, 비과세 제외)",
    "wife_gross": "아내 총급여(전년도, 비과세 제외)",
    "mortgage_borrower": "주담대 채무자(부부 합산 시 주 차주)",
    "credit_borrower": "신용대출 차주",
    "credit_use": "신용대출 반영",
    "credit_amount": "신용대출 금액(비우면 'DSR 기준 허용액' 자동 적용)",
    "credit_rate": "신용대출 금리(연)",
    "credit_method": "신용대출 상환 방식",
    "credit_draw_date": "신용대출 실행일",
    "credit_repay_date": "신용대출 상환일(비우면 계속 유지)",
    "sub_loan_use": "청약통장 담보대출 반영",
    "sub_principal_husband": "남편 청약통장 납입 원금",
    "sub_principal_wife": "아내 청약통장 납입 원금",
    "sub_ratio": "담보 인정 비율",
    "sub_amount": "대출 금액(비우면 한도 전액)",
    "sub_savings_rate": "청약통장 금리(연)",
    "sub_spread": "가산금리(연)",
    "sub_draw_date": "실행일",
    "sub_repay_date": "상환일(비우면 계속 유지)",
    "family_loan_rate": "가족 차입 이자율(연)",
    "family_loan_monthly_principal": "가족 차입 월 상환액(원금)",
    "family_loan_repay_start": "가족 차입 상환 시작일",
    # 명의·세금·비용
    "ownership": "명의 형태",
    "husband_share": "남편 지분",
    "first_home_tax_relief": "생애최초 취득세 감면 반영",
    "card_pay_tax": "취득 관련 세금 카드 납부",
    "card_months": "할부 개월 수",
    "card_paid_installments": "고객 부담 회차 수(앞에서부터)",
    "card_fee_rate": "카드 할부 수수료율(연)",
    "card_first_offset": "첫 회차 납부(잔금월 + N개월)",
    "lawyer_fee": "법무사 보수·등기 대행비",
    "registry_fee": "등기 신청 수수료·증지 등",
    "official_price_ratio": "공시가격(시가표준액) / 매매가 비율 가정",
    "bond_buy_rate": "국민주택채권 매입률",
    "bond_discount_rate": "국민주택채권 즉시매도 할인율",
    "loan_stamp_share": "대출 인지세 본인 부담 비율",
    "repair_default": "수리·인테리어 기본값",
    "appliances": "가전·가구",
    "moving": "이사비",
    "cleaning": "입주 청소",
    "temp_housing": "임시거주·보관 비용",
}

LISTING_COLS = {
    "id": "매물ID",
    "data_kind": "데이터 구분",
    "name": "단지명",
    "address": "주소·지역",
    "area_m2": "전용면적(㎡)",
    "unit": "동·층",
    "rooms": "방·욕실 수",
    "ask_price": "매도 호가(원)",
    "price": "검토 매수가(원)",
    "appraisal": "담보평가액(원, KB시세·감정가)",
    "mgmt_fee": "월 관리비(원)",
    "repair": "예상 수리비(원)",
    "move_in_available": "입주 가능일",
    "tenant": "임차인 퇴거 조건",
    "note": "비고",
}

LABELS.update(
    {
        "first_interest_adjustment": "첫 달 일할이자 조정액",
        "mortgage_round_unit": "역산 시 대출금 절사 단위",
        "husband_company_topup": "남편 회사 보전금(월)",
        "husband_company_topup_months": "남편 회사 보전 개월 수",
        "saving_in_spend_limit": "  └ 적금 납입을 월 가계지출 한도 판정에 포함?",
        "manual_tax_rate": "취득세율 수동 지정(중과 등, 선택)",
        "brokerage_rate": "중개보수 협의 요율(선택)",
    }
)


def norm(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def date_formula(formula: Any, cells: dict[str, str], named: dict[str, str]) -> str | None:
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    match = re.fullmatch(r"=([\w$]+)([+-]\d+)?", formula.replace(" ", ""))
    if not match:
        return None
    key = cells.get(match[1].replace("$", ""), named.get(match[1]))
    return key + (match[2] or "") if key in DATE_KEYS else None


def extract_excel(source: Path, extracted_at: date) -> dict[str, Any]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        values = load_workbook(source, data_only=True)
        formulas = load_workbook(source, data_only=False)
    try:
        settings_sheet = values["기본설정"]
        settings_formulas = formulas["기본설정"]
        rows = {
            settings_sheet.cell(row, 2).value: row
            for row in range(1, settings_sheet.max_row + 1)
            if isinstance(settings_sheet.cell(row, 2).value, str)
        }
        cells = {f"C{rows[label]}": key for key, label in LABELS.items() if label in rows}
        named = dict(ALIASES)
        for name, definition in formulas.defined_names.items():
            target = str(definition.attr_text).split("!")[-1].replace("$", "")
            if target in cells:
                named[name] = cells[target]
        missing = [key for key, label in LABELS.items() if label not in rows]
        settings: dict[str, Any] = {}
        statuses: dict[str, Any] = {}
        cached_formulas: dict[str, Any] = {}
        messages: list[str] = []

        def read_date(formula: Any, value: Any, location: str) -> Any:
            relative = date_formula(formula, cells, named)
            if relative:
                return relative
            if isinstance(formula, str) and formula.startswith("="):
                cached_formulas[location] = {"formula": formula, "cached": value}
                messages.append(f"지원하지 않는 날짜 수식: {location}; 캐시값 사용")
            return value

        for key, label in LABELS.items():
            if label not in rows:
                continue
            row = rows[label]
            value = norm(settings_sheet.cell(row, 5).value)
            raw = settings_formulas.cell(row, 3).value
            if key in DATE_KEYS:
                value = read_date(raw, value, f"기본설정!C{row}")
            settings[key] = value
            statuses[key] = norm(settings_sheet.cell(row, 8).value)
        settings["setting_statuses"] = statuses
        settings["comparison_area_m2"] = norm(values["조건별비교"]["B5"].value) or 84.9
        if values["기본설정"]["E66"].value == "수동":
            messages.append(
                "저축 수동 모드는 1차 범위에서 지원하지 않습니다. 자동 현금흐름으로 계산합니다."
            )
        cashflow = formulas["월별가계흐름"]
        manual = [
            cell.coordinate
            for row in cashflow.iter_rows(min_row=6, min_col=51, max_col=54)
            for cell in row
            if isinstance(cell.value, (int, float)) and cell.value != 0
        ]
        if manual:
            messages.append("월별 수동 조정은 지원하지 않습니다: " + ", ".join(manual))
        required = {"asset_base_date", "contract_date", "closing_date", "move_in_date"}
        if any(settings.get(key) in (None, 0, "") for key in required):
            raise ValueError("필수 날짜 캐시가 없습니다. 엑셀에서 재계산 후 저장하세요")
        events = []
        u, f = values["우리자금"], formulas["우리자금"]
        for row in range(6, 46):
            if not u.cell(row, 1).value:
                continue
            fields = {
                "id": 1,
                "data_kind": 2,
                "owner": 3,
                "type": 4,
                "name": 5,
                "value": 6,
                "cash_out_amount": 7,
                "cost": 8,
                "available_date": 10,
                "use_for_purchase": 11,
                "status": 12,
                "note": 14,
            }
            event = {key: norm(u.cell(row, col).value) for key, col in fields.items()}
            event["available_date"] = read_date(
                f.cell(row, 10).value, event["available_date"], f"우리자금!J{row}"
            )
            event["origin"] = "excel"
            events.append(event)
        parents: list[dict[str, Any]] = []
        for row in range(55, 95):
            if u.cell(row, 2).value != "순번":
                continue
            for child in range(row + 1, row + 13):
                if u.cell(child, 5).value is not None:
                    parents.append(
                        {
                            "id": f"P{len(parents) + 1:02}",
                            "to": u.cell(child, 3).value,
                            "kind": u.cell(child, 4).value,
                            "amount": norm(u.cell(child, 5).value),
                            "date": read_date(
                                f.cell(child, 6).value,
                                norm(u.cell(child, 6).value),
                                f"우리자금!F{child}",
                            ),
                            "status": u.cell(child, 7).value,
                            "note": u.cell(child, 8).value,
                            "origin": "excel",
                        }
                    )
            break
        listing_sheet = values["아파트매물"]
        headers = {
            listing_sheet.cell(5, col).value: col for col in range(1, listing_sheet.max_column + 1)
        }
        listings = [
            {
                key: norm(listing_sheet.cell(row, headers[label]).value)
                for key, label in LISTING_COLS.items()
                if label in headers
            }
            for row in range(6, 46)
            if listing_sheet.cell(row, 1).value
            and listing_sheet.cell(row, headers["데이터 구분"]).value == "실제"
        ]
        return {
            "_source": str(source),
            "_extracted_at": extracted_at.isoformat(),
            "_missing_labels": missing,
            "_import_warnings": messages,
            "_formula_cache": cached_formulas,
            "settings": settings,
            "fund_events": events,
            "parent_support": parents,
            "listings": listings,
            "selected_listing": values["결과요약"]["C6"].value,
        }
    finally:
        values.close()
        formulas.close()


def merge_import(
    existing: dict[str, Any], incoming: dict[str, Any], *, refresh_imported: bool = False
) -> dict[str, Any]:
    import copy

    merged = copy.deepcopy(existing)
    for key, value in incoming.items():
        if key not in {"settings", "fund_events", "parent_support", "listings", "bank_quotes"}:
            if (
                key not in merged
                or key.startswith("_import")
                or key in {"_extracted_at", "_source", "_formula_cache"}
            ):
                merged[key] = value
    app_keys = set(existing.get("app_setting_keys", []))
    merged["settings"] = {
        **existing.get("settings", {}),
        **{
            key: value
            for key, value in incoming["settings"].items()
            if key not in existing.get("settings", {}) or (refresh_imported and key not in app_keys)
        },
    }
    for section in ("fund_events", "parent_support", "listings"):
        old = copy.deepcopy(existing.get(section, []))
        if section == "parent_support":
            for i, item in enumerate(old, 1):
                if not item.get("id"):
                    item["id"] = f"P{i:02}"
        by_id = {item.get("id"): item for item in old}
        for item in incoming[section]:
            previous = by_id.get(item.get("id"))
            if previous is None:
                old.append(copy.deepcopy(item))
            elif refresh_imported and previous.get("origin") != "app":
                previous.update(copy.deepcopy(item))
        merged[section] = old
    if existing:
        merged["_import_warnings"] = list(incoming.get("_import_warnings", [])) + [
            "엑셀은 가져오기 원본입니다. "
            + (
                "명시적 갱신: 기존 Excel 출처 항목·설정을 갱신했습니다. 앱 항목은 유지합니다."
                if refresh_imported
                else "기존 앱 데이터가 우선이며, 중복 ID·기존 설정을 보존하고 새 항목만 추가했습니다."
            )
        ]
    for qi, quote in enumerate(merged.get("bank_quotes", []), 1):
        quote.setdefault("id", f"q{qi}")
        for oi, option in enumerate(quote["options"], 1):
            if not option.get("id"):
                option["id"] = f"{quote['id']}-o{oi}"
    if not merged["settings"].get("mortgage_quote_option"):
        selections: dict[str, str] = {}
        for quote in sorted(merged.get("bank_quotes", []), key=lambda q: q["date"], reverse=True):
            for option in quote["options"]:
                selections.setdefault(option["method"], option["id"])
        merged["settings"]["mortgage_quote_option"] = selections
    return merged


def import_excel(
    source: Path, destination: Path, extracted_at: date, *, refresh_imported: bool = False
) -> tuple[dict[str, Any], Path | None]:
    existing = read_json(destination) if destination.exists() else {}
    merged = merge_import(
        existing, extract_excel(source, extracted_at), refresh_imported=refresh_imported
    )
    return merged, save_json(destination, merged)
